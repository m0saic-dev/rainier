#!/usr/bin/env python3
"""Build kits/marketing-review.xlsx: the post log and review dashboard.

From the repo root:

    python3 -m venv .venv
    .venv/bin/pip install openpyxl==3.1.5
    .venv/bin/python kits/build-marketing-review.py             # writes kits/marketing-review.xlsx
    .venv/bin/python kits/build-marketing-review.py --out x.xlsx

    .venv/bin/pip install formulas==1.3.4                        # optional, for --verify
    .venv/bin/python kits/build-marketing-review.py --verify     # evaluates every formula

The workbook holds made-up sample rows only. Output is deterministic: fixed sample
data, fixed document properties and fixed zip entry times, so the same openpyxl and
zlib rebuild it byte for byte.

Portability (Excel for Mac and Windows, Apple Numbers, Google Sheets): plain absolute
ranges, no structured table references, no dynamic-array functions; WEEKNUM(date, 2),
SUMIFS / COUNTIFS / AVERAGEIFS / RANK / IFERROR / INDEX / MATCH / LOOKUP / HYPERLINK
only. The Posts rank is two COUNTIFS (ties go to the row logged first).

A row counts as a post once it has a date: every Review aggregate carries the same
condition, and the Posts rank skips undated rows, so the numbers agree. The condition
reads Posts N (Week), a plain number exactly when A holds a date: ">0" on a number column
means the same in Excel, Numbers and Sheets, where "<>" on the date column leans on each
app's blank handling (and the formulas package used by --verify reads "<>" as '< ">"').

No formula cell carries a cached value (openpyxl writes none), so a viewer that does
not calculate (Quick Look, a phone's file preview, Excel's Protected View) shows the
formula cells and charts empty. The README sheet says so.
"""

import argparse
import datetime
import io
import math
import os
import sys
import tempfile
import zipfile

import openpyxl
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.data_source import AxDataSource, StrRef
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.writer.excel import ExcelWriter

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_OUT = os.path.join(HERE, "marketing-review.xlsx")
BUILT_WITH = "3.1.5"      # openpyxl version the committed file was built with
VERIFIED_WITH = "1.3.4"   # formulas version --verify was written against (it reads its solution keys)

FIXED_TIME = datetime.datetime(2026, 9, 27, 12, 0, 0)
ZIP_TIME = (2026, 9, 27, 12, 0, 0)

LAST = 1000       # Posts: formulas and formats run to this row (999 posts)
CAL_LAST = 500    # Calendar: validation and formats run to this row

# ---------------------------------------------------------------- palette ---
FONT = "Arial"
INK = "1F1F1D"
INK2 = "52514E"
MUTED = "8A8983"
RULE = "D9D8D4"
HEAD_BG = "2F2F2C"     # input column headers
DERIVED_BG = "6B6A65"  # formula column headers
DERIVED_FILL = "F2F2EF"
ACCENT = "2A78D6"      # the one chart colour

# ------------------------------------------------------------------ lists ---
PLATFORMS = ["Instagram", "TikTok", "YouTube"]
SECTIONS = ["intro", "verse 1", "pre-chorus", "chorus", "verse 2", "bridge", "outro"]
YES_NO = ["Yes", "No"]
STATUSES = ["Idea", "Cut", "Edited", "Scheduled", "Posted"]
BUCKETS = [(0, "0–15 s"), (15, "15–25 s"), (25, "25–40 s"), (40, "40+ s")]

LIST_PLATFORM = "Lists!$A$2:$A$%d" % (1 + len(PLATFORMS))
LIST_SECTION = "Lists!$B$2:$B$%d" % (1 + len(SECTIONS))
LIST_YES_NO = "Lists!$C$2:$C$%d" % (1 + len(YES_NO))
LIST_STATUS = "Lists!$D$2:$D$%d" % (1 + len(STATUSES))
LIST_BUCKET_FROM = "Lists!$F$2:$F$%d" % (1 + len(BUCKETS))
LIST_BUCKET_LABEL = "Lists!$G$2:$G$%d" % (1 + len(BUCKETS))

OTHER_PLATFORM = "Other or no platform"  # Review row for dated posts outside the platform list

# ------------------------------------------------------------ sample data ---
# 24 made-up posts, Aug-Sep 2026. Clip = the Take Cutter's file name
# (<take>_<MMmSSs>-<MMmSSs>, source timestamps). Views None = not in yet.
D = datetime.date
SAMPLE_POSTS = [
    # date, platform, clip, section, length s, views, likes, comments, shares, saves, trial, notes
    (D(2026, 8, 4), "TikTok", "verse-1_03m10s-03m28s", "verse 1", 18, 2140, 161, 9, 12, 21, "No", "First post of the song"),
    (D(2026, 8, 4), "Instagram", "verse-1_03m10s-03m28s", "verse 1", 18, 1630, 118, 6, 9, 17, "No", ""),
    (D(2026, 8, 5), "YouTube", "verse-1_03m10s-03m28s", "verse 1", 18, 880, 41, 2, 3, None, "No", "Title 64 characters"),
    (D(2026, 8, 11), "TikTok", "chorus-2_01m22s-01m41s", "chorus", 19, 8420, 702, 38, 96, 131, "No", "Posted at 7 pm"),
    (D(2026, 8, 11), "Instagram", "chorus-2_01m22s-01m41s", "chorus", 19, 5210, 433, 21, 58, 88, "No", ""),
    (D(2026, 8, 12), "YouTube", "chorus-2_01m22s-01m41s", "chorus", 19, 3960, 214, 11, 19, None, "No", ""),
    (D(2026, 8, 18), "TikTok", "intro_00m04s-00m17s", "intro", 13, 1210, 84, 3, 5, 9, "No", ""),
    (D(2026, 8, 18), "Instagram", "intro_00m04s-00m17s", "intro", 13, 940, 61, 2, 4, 8, "No", ""),
    (D(2026, 8, 19), "YouTube", "intro_00m04s-00m17s", "intro", 13, 610, 22, 1, 1, None, "No", ""),
    (D(2026, 8, 25), "Instagram", "bridge_27m40s-28m13s", "bridge", 33, 3870, 297, 17, 31, 52, "No", "Main post, two trial reels the same day"),
    (D(2026, 8, 25), "Instagram", "chorus-1_12m04s-12m16s", "chorus", 12, 2650, 171, 5, 22, 30, "Yes", "Trial reel A"),
    (D(2026, 8, 25), "Instagram", "pre-chorus_08m51s-09m13s", "pre-chorus", 22, 4480, 356, 14, 47, 61, "Yes", "Trial reel B"),
    (D(2026, 8, 25), "TikTok", "bridge_27m40s-28m13s", "bridge", 33, 5930, 468, 26, 55, 83, "No", ""),
    (D(2026, 8, 26), "YouTube", "bridge_27m40s-28m13s", "bridge", 33, 9870, 512, 29, 41, None, "No", "Picked up by the Shorts feed"),
    (D(2026, 9, 1), "TikTok", "verse-2_19m30s-19m52s", "verse 2", 22, 3350, 247, 12, 20, 34, "No", ""),
    (D(2026, 9, 1), "Instagram", "verse-2_19m30s-19m52s", "verse 2", 22, 2080, 149, 7, 13, 22, "No", ""),
    (D(2026, 9, 8), "TikTok", "outro_41m02s-41m49s", "outro", 47, 1480, 92, 4, 6, 11, "No", ""),
    (D(2026, 9, 8), "Instagram", "outro_41m02s-41m49s", "outro", 47, 1210, 70, 3, 5, 9, "No", "Same views as the intro TikTok"),
    (D(2026, 9, 9), "YouTube", "outro_41m02s-41m49s", "outro", 47, 690, 28, 1, 2, None, "No", ""),
    (D(2026, 9, 15), "TikTok", "chorus-3_33m15s-33m39s", "chorus", 24, 12650, 1104, 57, 163, 214, "No", "Best hook so far"),
    (D(2026, 9, 15), "Instagram", "chorus-3_33m15s-33m39s", "chorus", 24, 7340, 618, 31, 84, 117, "No", ""),
    (D(2026, 9, 22), "TikTok", "pre-chorus_08m51s-09m13s", "pre-chorus", 22, 4120, 318, 15, 36, 49, "No", "Trial reel B, now a main post"),
    (D(2026, 9, 22), "Instagram", "pre-chorus_08m51s-09m13s", "pre-chorus", 22, 2960, 224, 10, 27, 38, "No", ""),
    (D(2026, 9, 24), "YouTube", "pre-chorus_08m51s-09m13s", "pre-chorus", 22, None, None, None, None, None, "No", "Fill in the numbers after a week"),
]

SAMPLE_CALENDAR = [
    # planned date, song, section, take file, platform(s), status, posted link (post number or None)
    (D(2026, 9, 22), "Song Title", "pre-chorus", "pre-chorus_08m51s-09m13s.mp4", "TikTok, Instagram, YouTube", "Posted", 22),
    (D(2026, 9, 29), "Song Title", "chorus", "chorus-4_36m02s-36m23s.mp4", "TikTok, Instagram, YouTube", "Scheduled", None),
    (D(2026, 9, 29), "Song Title", "verse 1", "verse-1b_05m41s-06m00s.mp4", "Instagram trial reel", "Edited", None),
    (D(2026, 9, 29), "Song Title", "bridge", "bridge-2_29m10s-29m38s.mp4", "Instagram trial reel", "Edited", None),
    (D(2026, 10, 6), "Second Song", "chorus", "chorus_02m15s-02m36s.mp4", "TikTok, Instagram, YouTube", "Cut", None),
    (D(2026, 10, 6), "Second Song", "intro", "", "TikTok, YouTube", "Idea", None),
]


def post_link(platform, n):
    """Placeholder link for sample post n (1-based)."""
    if platform == "TikTok":
        return "https://www.tiktok.com/@yourhandle/video/%019d" % n
    if platform == "Instagram":
        return "https://www.instagram.com/reel/placeholder%02d/" % n
    return "https://www.youtube.com/shorts/placeholder%02d" % n


# ------------------------------------------------------------- helpers ---
def font(size=10, bold=False, italic=False, color=INK, underline=None):
    return Font(name=FONT, size=size, bold=bold, italic=italic, color=color, underline=underline)


def fill(rgb):
    return PatternFill(fill_type="solid", start_color=rgb, end_color=rgb)


RULE_SIDE = Side(style="thin", color=RULE)
INK_SIDE = Side(style="thin", color=INK2)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
HEAD_ALIGN = Alignment(wrap_text=True, vertical="center")


def rng(col):
    """Absolute Posts range for one column, rows 2..LAST."""
    return "Posts!$%s$2:$%s$%d" % (col, col, LAST)


def dated():
    """The criteria pair every Review SUMIFS / COUNTIFS / AVERAGEIFS carries: a row counts
    as a post once it has a date. Posts N (Week) is a number exactly when A holds a date."""
    return '%s,">0"' % rng("N")


def set_widths(ws, widths):
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def list_validation(ws, source, cells, error):
    dv = DataValidation(type="list", formula1=source, allow_blank=True)
    dv.showErrorMessage = True
    dv.errorTitle = "Pick from the list"
    dv.error = error
    dv.add(cells)
    ws.add_data_validation(dv)


RENAME = " To change the choices, rename one on the Lists sheet."
PICK_PLATFORM = "Pick a platform from the list." + RENAME
PICK_SECTION = "Pick a song section from the list." + RENAME
PICK_STATUS = "Pick a status from the list." + RENAME


def date_validation(ws, cells):
    dv = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True)
    dv.showErrorMessage = True
    dv.errorTitle = "Type a date"
    dv.error = "Type a date, like 2026-09-27."
    dv.add(cells)
    ws.add_data_validation(dv)


def number_validation(ws, cells):
    dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
    dv.showErrorMessage = True
    dv.errorTitle = "Type a number"
    dv.error = "Type a number, 0 or more (no text, no 'k')."
    dv.add(cells)
    ws.add_data_validation(dv)


# --------------------------------------------------------------- Lists ---
def build_lists(ws):
    heads = {"A": "Platform", "B": "Section", "C": "Yes / No", "D": "Status", "F": "From (s)", "G": "Length bucket"}
    for col, text in heads.items():
        cell = ws["%s1" % col]
        cell.value = text
        cell.font = font(bold=True, color="FFFFFF")
        cell.fill = fill(HEAD_BG)
        cell.alignment = HEAD_ALIGN
    for i, v in enumerate(PLATFORMS):
        ws.cell(2 + i, 1, v)
    for i, v in enumerate(SECTIONS):
        ws.cell(2 + i, 2, v)
    for i, v in enumerate(YES_NO):
        ws.cell(2 + i, 3, v)
    for i, v in enumerate(STATUSES):
        ws.cell(2 + i, 4, v)
    for i, (start, label) in enumerate(BUCKETS):
        ws.cell(2 + i, 6, start).number_format = "0"
        ws.cell(2 + i, 7, label)
    for row in ws.iter_rows(min_row=2, max_row=1 + len(SECTIONS), max_col=7):
        for cell in row:
            cell.font = font()
    notes = [
        "These lists feed the dropdowns on Posts and Calendar, and the Review sheet.",
        "Each list has a fixed number of rows: to change a choice, rename it. A row added below a list is not picked up.",
        "Rename an entry here and the dropdowns follow; posts already logged keep the old name (Find and Replace on Posts).",
        "Review shows the platforms, sections and length buckets listed here; a platform not listed counts as Other.",
        "The trial-reel rows on Review look for the names Instagram and Yes, so keep those two as they are.",
        "Length buckets: a clip goes in the last bucket whose From (s) it reaches.",
        "Keep From (s) in increasing order.",
    ]
    for i, text in enumerate(notes):
        ws.cell(1 + i, 9, text).font = font(color=INK2, italic=True)
    set_widths(ws, {"A": 13, "B": 13, "C": 10, "D": 12, "E": 3, "F": 10, "G": 14, "H": 3, "I": 100})
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"


# --------------------------------------------------------------- Posts ---
POST_COLUMNS = [
    # letter, header, width, number format, derived
    ("A", "Date", 12, "d mmm yyyy", False),
    ("B", "Platform", 11, "General", False),
    ("C", "Clip", 27, "General", False),
    ("D", "Section", 11, "General", False),
    ("E", "Link", 40, "General", False),
    ("F", "Length (s)", 9, "0", False),
    ("G", "Views", 10, "#,##0", False),
    ("H", "Likes", 9, "#,##0", False),
    ("I", "Comments", 10, "#,##0", False),
    ("J", "Shares", 9, "#,##0", False),
    ("K", "Saves", 9, "#,##0", False),
    ("L", "Trial reel", 8, "General", False),
    ("M", "Notes", 34, "General", False),
    ("N", "Week", 7, "0", True),
    ("O", "Length bucket", 10, "General", True),
    ("P", "Engagement", 11, "#,##0", True),
    ("Q", "Engagement rate", 11, "0.0%", True),
    ("R", "Rank", 7, "0", True),
]


def posts_formulas(r):
    return {
        "N": '=IF(A{r}="","",WEEKNUM(A{r},2))',
        "O": '=IF(OR(A{r}="",F{r}=""),"",LOOKUP(F{r},%s,%s))' % (LIST_BUCKET_FROM, LIST_BUCKET_LABEL),
        "P": '=IF(A{r}="","",SUM(H{r}:K{r}))',
        "Q": '=IF(A{r}="","",IFERROR(P{r}/G{r},""))',
        # dated rows only (N is a number exactly when A holds a date): the dated rows with
        # more views, plus the dated rows up to this one with the same views (ties go to the
        # row logged first)
        "R": ('=IF(OR(A{r}="",G{r}=""),"",COUNTIFS($N$2:$N$%d,">0",$G$2:$G$%d,">"&G{r})'
              '+COUNTIFS($N$2:N{r},">0",$G$2:G{r},G{r}))') % (LAST, LAST),
    }


def build_posts(ws, posts):
    for letter, header, width, fmt, derived in POST_COLUMNS:
        cell = ws["%s1" % letter]
        cell.value = header
        cell.font = font(bold=True, color="FFFFFF")
        cell.fill = fill(DERIVED_BG if derived else HEAD_BG)
        cell.alignment = HEAD_ALIGN
        ws.column_dimensions[letter].width = width
    ws.row_dimensions[1].height = 30

    link_font = font(color=ACCENT, underline="single")
    for i, post in enumerate(posts):
        r = 2 + i
        date, platform, clip, section, length, views, likes, comments, shares, saves, trial, notes = post
        values = [date, platform, clip, section, post_link(platform, i + 1), length,
                  views, likes, comments, shares, saves, trial, notes or None]
        for c, v in enumerate(values):
            if v is not None:
                ws.cell(r, 1 + c, v)
        ws.cell(r, 5).hyperlink = post_link(platform, i + 1)

    for r in range(2, LAST + 1):
        formulas_row = {k: v.format(r=r) for k, v in posts_formulas(r).items()}
        for letter, _header, _width, fmt, derived in POST_COLUMNS:
            cell = ws["%s%d" % (letter, r)]
            if derived:
                cell.value = formulas_row[letter]
                cell.fill = fill(DERIVED_FILL)
                cell.font = font(color=INK2)
            elif letter == "E" and r <= 1 + len(posts):
                cell.font = link_font
            else:
                cell.font = font()
            cell.number_format = fmt

    last = str(LAST)
    date_validation(ws, "A2:A" + last)
    list_validation(ws, LIST_PLATFORM, "B2:B" + last, PICK_PLATFORM)
    list_validation(ws, LIST_SECTION, "D2:D" + last, PICK_SECTION)
    number_validation(ws, "F2:K" + last)
    list_validation(ws, LIST_YES_NO, "L2:L" + last, "Pick Yes or No.")
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = "A1:R" + last


# ------------------------------------------------------------ Calendar ---
CAL_COLUMNS = [
    ("A", "Planned date", 13, "d mmm yyyy"),
    ("B", "Song", 18, "General"),
    ("C", "Section", 11, "General"),
    ("D", "Take file", 32, "General"),
    ("E", "Platform(s)", 26, "General"),
    ("F", "Status", 11, "General"),
    ("G", "Posted link", 46, "General"),
]


def build_calendar(ws, calendar, posts):
    for letter, header, width, _fmt in CAL_COLUMNS:
        cell = ws["%s1" % letter]
        cell.value = header
        cell.font = font(bold=True, color="FFFFFF")
        cell.fill = fill(HEAD_BG)
        cell.alignment = HEAD_ALIGN
        ws.column_dimensions[letter].width = width
    ws.row_dimensions[1].height = 30
    link_font = font(color=ACCENT, underline="single")
    for r in range(2, CAL_LAST + 1):
        for letter, _header, _width, fmt in CAL_COLUMNS:
            cell = ws["%s%d" % (letter, r)]
            cell.font = font()
            cell.number_format = fmt
    for i, (date, song, section, take, platforms, status, post_no) in enumerate(calendar):
        r = 2 + i
        for c, v in enumerate([date, song, section, take or None, platforms, status]):
            if v is not None:
                ws.cell(r, 1 + c, v)
        if post_no is not None:
            platform = posts[post_no - 1][1]
            link = ws.cell(r, 7, post_link(platform, post_no))
            link.hyperlink = post_link(platform, post_no)
            link.font = link_font
    last = str(CAL_LAST)
    date_validation(ws, "A2:A" + last)
    list_validation(ws, LIST_SECTION, "C2:C" + last, PICK_SECTION)
    list_validation(ws, LIST_STATUS, "F2:F" + last, PICK_STATUS)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:G" + last


# -------------------------------------------------------------- Review ---
CHART_W = 15.5   # cm
CHART_H = 6.8    # cm, about 13 default rows
CHART_ROWS = 14


def style_chart(ch, title, numfmt="#,##0", horizontal=False):
    ch.title = title
    ch.legend = None
    ch.width = CHART_W
    ch.height = CHART_H
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.numFmt = numfmt
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=RULE))
    ch.y_axis.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
    ch.x_axis.spPr = GraphicalProperties(ln=LineProperties(solidFill=RULE))
    ch.x_axis.axPos = "l" if horizontal else "b"  # category axis
    ch.y_axis.axPos = "b" if horizontal else "l"  # value axis


def bar_chart(ws, direction, title, data_col, cat_col, head_row, first, last, anchor):
    ch = BarChart()
    ch.type = direction  # "col" = vertical bars, "bar" = horizontal bars
    ch.grouping = "clustered"
    ch.gapWidth = 80
    ch.add_data(Reference(ws, min_col=data_col, min_row=head_row, max_row=last), titles_from_data=True)
    style_chart(ch, title, horizontal=(direction == "bar"))
    series = ch.series[0]
    # text categories as a string reference, so every app labels the bars with the names
    series.cat = AxDataSource(strRef=StrRef(f=str(Reference(ws, min_col=cat_col, min_row=first, max_row=last))))
    series.graphicalProperties.solidFill = ACCENT
    series.graphicalProperties.line.noFill = True
    if direction == "bar":
        # list order top to bottom, value axis kept at the bottom
        ch.x_axis.scaling.orientation = "maxMin"
        ch.y_axis.crosses = "max"
    ws.add_chart(ch, anchor)


def line_chart(ws, title, data_col, cat_col, head_row, first, last, anchor):
    ch = LineChart()
    ch.add_data(Reference(ws, min_col=data_col, min_row=head_row, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=cat_col, min_row=first, max_row=last))
    style_chart(ch, title)
    ch.x_axis.number_format = "d mmm"
    series = ch.series[0]
    series.smooth = False
    series.graphicalProperties.line.solidFill = ACCENT
    series.graphicalProperties.line.width = 19050  # 1.5 pt
    series.marker.symbol = "circle"
    series.marker.size = 6
    series.marker.graphicalProperties = GraphicalProperties(solidFill=ACCENT, ln=LineProperties(solidFill=ACCENT))
    ws.add_chart(ch, anchor)


class Review:
    """Writes the Review sheet top to bottom and records where each checked value lives."""

    def __init__(self, ws):
        self.ws = ws
        self.cells = {}  # check name -> coordinate, read by --verify

    def put(self, row, col, value, fmt=None, name=None, **font_kw):
        cell = self.ws.cell(row, col, value)
        cell.font = font(**font_kw)
        if fmt:
            cell.number_format = fmt
        if name:
            self.cells[name] = cell.coordinate
        return cell

    def title(self, row, text, note=None):
        self.put(row, 2, text, size=12, bold=True)
        if note:
            self.put(row + 1, 2, note, color=INK2, italic=True)
            return row + 2
        return row + 1

    def header(self, row, first_col, labels):
        for i, text in enumerate(labels):
            cell = self.put(row, first_col + i, text, bold=True, color=INK2)
            cell.border = Border(bottom=INK_SIDE)
            cell.alignment = Alignment(horizontal="left" if i == 0 else "right", wrap_text=True, vertical="bottom")
        self.ws.row_dimensions[row].height = 26

    def rule_row(self, row, first_col, last_col):
        for col in range(first_col, last_col + 1):
            self.ws.cell(row, col).border = Border(top=INK_SIDE)


def build_review(ws):
    rv = Review(ws)
    set_widths(ws, {"A": 4, "B": 34, "C": 12, "D": 12, "E": 12, "F": 12, "G": 12, "H": 3})
    ws.sheet_view.showGridLines = False
    rv.put(1, 2, "Marketing review", size=18, bold=True)
    rv.put(2, 2, "Worked out from the Posts sheet as you log. Nothing to type on this page.", color=INK2)

    G, P, DATED = rng("G"), rng("P"), dated()

    # At a glance (formulas for the three "best" rows are filled in at the end)
    r = rv.title(4, "At a glance")
    kpi_rows = {}
    kpis = [
        ("posts", "Posts logged", "=COUNTIFS(%s)" % DATED, "#,##0"),
        ("views", "Total views", "=SUMIFS(%s,%s)" % (G, DATED), "#,##0"),
        ("avg", "Average views per post", '=IFERROR(AVERAGEIFS(%s,%s),"")' % (G, DATED), "#,##0"),
        ("rate", "Engagement rate", '=IFERROR(SUMIFS(%s,%s,%s,">0")/SUMIFS(%s,%s),"")' % (P, DATED, G, G, DATED),
         "0.0%"),
        ("best_platform", "Best platform (average views)", None, None),
        ("best_length", "Best length (average views)", None, None),
        ("best_section", "Best song section (average views)", None, None),
    ]
    for key, label, formula, fmt in kpis:
        rv.put(r, 2, label, color=INK2)
        cell = rv.put(r, 3, formula, fmt, name="kpi." + key, size=11, bold=True)
        cell.alignment = Alignment(horizontal="right")
        ws.cell(r, 2).border = Border(bottom=RULE_SIDE)
        cell.border = Border(bottom=RULE_SIDE)
        kpi_rows[key] = r
        r += 1
    r += 1

    # Views by platform (bar chart)
    top = r
    r = rv.title(r, "Views by platform")
    rv.header(r, 2, ["Platform", "Posts", "Views", "Avg views", "Engagement rate", "Share of views"])
    head = r
    r += 1
    plat_first = r
    B = rng("B")
    all_row = r + len(PLATFORMS) + 1  # the platforms, the Other row, then All platforms
    for i, name in enumerate(PLATFORMS):
        rv.put(r, 2, "=Lists!$A$%d" % (2 + i))
        rv.put(r, 3, "=COUNTIFS(%s,%s,$B%d)" % (DATED, B, r), "#,##0", "platform.%s.posts" % name)
        rv.put(r, 4, "=SUMIFS(%s,%s,%s,$B%d)" % (G, DATED, B, r), "#,##0", "platform.%s.views" % name)
        rv.put(r, 5, '=IFERROR(AVERAGEIFS(%s,%s,%s,$B%d),"")' % (G, DATED, B, r), "#,##0", "platform.%s.avg" % name)
        rv.put(r, 6, '=IFERROR(SUMIFS(%s,%s,%s,$B%d,%s,">0")/D%d,"")' % (P, DATED, B, r, G, r), "0.0%",
               "platform.%s.rate" % name)
        rv.put(r, 7, '=IFERROR(D%d/$D$%d,"")' % (r, all_row), "0.0%", "platform.%s.share" % name)
        r += 1
    plat_last = r - 1
    # dated posts whose platform is empty or not on the list, so the column adds up; blank when none
    rv.put(r, 2, '=IF(C{r}="","","{label}")'.format(r=r, label=OTHER_PLATFORM), name="platform.other.label",
           color=INK2)
    rv.put(r, 3, '=IF(COUNTIFS({d})=SUM(C{a}:C{b}),"",COUNTIFS({d})-SUM(C{a}:C{b}))'.format(
        d=DATED, a=plat_first, b=plat_last), "#,##0", "platform.other.posts", color=INK2)
    rv.put(r, 4, '=IF(C{r}="","",SUMIFS({G},{d})-SUM(D{a}:D{b}))'.format(
        r=r, G=G, d=DATED, a=plat_first, b=plat_last), "#,##0", "platform.other.views", color=INK2)
    rv.put(r, 7, '=IF(C{r}="","",IFERROR(D{r}/$D${t},""))'.format(r=r, t=all_row), "0.0%",
           "platform.other.share", color=INK2)
    r += 1
    rv.rule_row(r, 2, 7)
    rv.put(r, 2, "All platforms", bold=True)
    rv.put(r, 3, "=SUM(C%d:C%d)" % (plat_first, r - 1), "#,##0", "platform.all.posts", bold=True)
    rv.put(r, 4, "=SUM(D%d:D%d)" % (plat_first, r - 1), "#,##0", "platform.all.views", bold=True)
    rv.put(r, 5, '=IFERROR(AVERAGEIFS(%s,%s),"")' % (G, DATED), "#,##0", "platform.all.avg", bold=True)
    rv.put(r, 6, "=C%d" % kpi_rows["rate"], "0.0%", "platform.all.rate", bold=True)
    rv.put(r, 7, '=IFERROR(D%d/$D$%d,"")' % (r, all_row), "0.0%", "platform.all.share", bold=True)
    assert r == all_row
    bar_chart(ws, "col", "Views by platform", 4, 2, head, plat_first, plat_last, "I%d" % top)
    r = max(r + 2, top + CHART_ROWS + 1)

    # Average views by length bucket (column chart)
    top = r
    r = rv.title(r, "Average views by length",
                 "A clip of exactly 15 s counts as 15–25 s. The cut-offs live on the Lists sheet.")
    rv.header(r, 2, ["Length", "Posts", "Avg views", "Engagement rate"])
    head = r
    r += 1
    len_first = r
    O = rng("O")
    for i, (_start, label) in enumerate(BUCKETS):
        rv.put(r, 2, "=Lists!$G$%d" % (2 + i))
        rv.put(r, 3, "=COUNTIFS(%s,%s,$B%d)" % (DATED, O, r), "#,##0", "length.%d.posts" % i)
        rv.put(r, 4, '=IFERROR(AVERAGEIFS(%s,%s,%s,$B%d),"")' % (G, DATED, O, r), "#,##0", "length.%d.avg" % i)
        rv.put(r, 5, '=IFERROR(SUMIFS(%s,%s,%s,$B%d,%s,">0")/SUMIFS(%s,%s,%s,$B%d),"")' % (
            P, DATED, O, r, G, G, DATED, O, r), "0.0%", "length.%d.rate" % i)
        r += 1
    len_last = r - 1
    bar_chart(ws, "col", "Average views by length", 4, 2, head, len_first, len_last, "I%d" % top)
    r = max(r + 1, top + CHART_ROWS + 1)

    # Top 10 posts by views
    r = rv.title(r, "Top 10 clips", "Most views first. Ties go to the post logged first.")
    rv.header(r, 1, ["#", "Clip", "Platform", "Section", "Length (s)", "Views", "Link"])
    ws.cell(r, 7).alignment = Alignment(horizontal="left", vertical="bottom")
    r += 1
    R = rng("R")
    for k in range(1, 11):
        m = "MATCH($A%d,%s,0)" % (r, R)
        rv.put(r, 1, k, color=INK2)
        rv.put(r, 2, '=IFERROR(INDEX(%s,%s)&"","")' % (rng("C"), m), name="top.%d.clip" % k)
        rv.put(r, 3, '=IFERROR(INDEX(%s,%s)&"","")' % (rng("B"), m), name="top.%d.platform" % k).alignment = Alignment(horizontal="right")
        rv.put(r, 4, '=IFERROR(INDEX(%s,%s)&"","")' % (rng("D"), m), name="top.%d.section" % k).alignment = Alignment(horizontal="right")
        rv.put(r, 5, '=IFERROR(IF(INDEX({F},{m})="","",INDEX({F},{m})),"")'.format(F=rng("F"), m=m), "0",
               "top.%d.length" % k)
        rv.put(r, 6, '=IFERROR(INDEX(%s,%s),"")' % (G, m), "#,##0", "top.%d.views" % k)
        rv.put(r, 7, '=IFERROR(IF(INDEX({E},{m})="","",HYPERLINK(INDEX({E},{m}))),"")'.format(E=rng("E"), m=m),
               name="top.%d.link" % k, color=ACCENT, underline="single")
        for col in range(1, 8):
            ws.cell(r, col).border = Border(bottom=RULE_SIDE)
        r += 1
    r += 1

    # 12-week trend (line chart)
    top = r
    r = rv.title(r, "12-week trend", "The 12 weeks up to your latest post. Weeks start on Monday.")
    rv.header(r, 2, ["Week of", "Posts", "Views", "Avg views"])
    head = r
    r += 1
    trend_first = r
    A = rng("A")
    for k in range(12):
        if k == 0:
            start = '=IF(COUNT({A})=0,"",MAX({A})-WEEKDAY(MAX({A}),2)+1-77)'.format(A=A)
        else:
            start = '=IF(B{p}="","",B{p}+7)'.format(p=r - 1)
        win = '{A},">="&B{r},{A},"<"&(B{r}+7)'.format(A=A, r=r)
        rv.put(r, 2, start, "d mmm yyyy", "trend.%d.start" % k).alignment = Alignment(horizontal="left")
        rv.put(r, 3, '=IF(B{r}="","",COUNTIFS({w}))'.format(r=r, w=win), "#,##0", "trend.%d.posts" % k)
        rv.put(r, 4, '=IF(B{r}="","",SUMIFS({G},{w}))'.format(r=r, G=G, w=win), "#,##0", "trend.%d.views" % k)
        rv.put(r, 5, '=IF(B{r}="","",IFERROR(AVERAGEIFS({G},{w}),""))'.format(r=r, G=G, w=win), "#,##0",
               "trend.%d.avg" % k)
        r += 1
    trend_last = r - 1
    line_chart(ws, "Views per week", 4, 2, head, trend_first, trend_last, "I%d" % top)
    r = max(r + 1, top + CHART_ROWS + 1)

    # Trial reels against main posts
    r = rv.title(r, "Trial reels against main posts",
                 "Trial reels are an Instagram feature, so compare them with Instagram main posts.")
    rv.header(r, 2, ["", "Posts", "Views", "Avg views", "Engagement rate"])
    r += 1
    L = rng("L")
    rows = [
        # blank Trial reel counts as a main post ("<>Yes" matches empty cells)
        ("main", "Main posts, all platforms",
         '=COUNTIFS({d},{L},"<>Yes")',
         '=SUMIFS({G},{d},{L},"<>Yes")',
         '=IFERROR(AVERAGEIFS({G},{d},{L},"<>Yes"),"")',
         '=IFERROR(SUMIFS({P},{d},{L},"<>Yes",{G},">0")/D{r},"")'),
        ("ig_main", "Instagram main posts",
         '=COUNTIFS({d},{B},"Instagram",{L},"<>Yes")',
         '=SUMIFS({G},{d},{B},"Instagram",{L},"<>Yes")',
         '=IFERROR(AVERAGEIFS({G},{d},{B},"Instagram",{L},"<>Yes"),"")',
         '=IFERROR(SUMIFS({P},{d},{B},"Instagram",{L},"<>Yes",{G},">0")/D{r},"")'),
        ("trial", "Trial reels",
         '=COUNTIFS({d},{L},"Yes")',
         '=SUMIFS({G},{d},{L},"Yes")',
         '=IFERROR(AVERAGEIFS({G},{d},{L},"Yes"),"")',
         '=IFERROR(SUMIFS({P},{d},{L},"Yes",{G},">0")/D{r},"")'),
    ]
    for key, label, f_posts, f_views, f_avg, f_rate in rows:
        env = dict(d=DATED, B=B, G=G, L=L, P=P, r=r)
        rv.put(r, 2, label)
        rv.put(r, 3, f_posts.format(**env), "#,##0", "trial.%s.posts" % key)
        rv.put(r, 4, f_views.format(**env), "#,##0", "trial.%s.views" % key)
        rv.put(r, 5, f_avg.format(**env), "#,##0", "trial.%s.avg" % key)
        rv.put(r, 6, f_rate.format(**env), "0.0%", "trial.%s.rate" % key)
        for col in range(2, 7):
            ws.cell(r, col).border = Border(bottom=RULE_SIDE)
        r += 1
    r += 1

    # Best song sections (horizontal bar chart)
    top = r
    r = rv.title(r, "Best song sections", "In song order; Rank 1 has the most views per post.")
    rv.header(r, 2, ["Section", "Posts", "Views", "Avg views", "Rank"])
    head = r
    r += 1
    sec_first = r
    sec_last = r + len(SECTIONS) - 1
    Dcol = rng("D")
    for i, name in enumerate(SECTIONS):
        rv.put(r, 2, "=Lists!$B$%d" % (2 + i))
        rv.put(r, 3, "=COUNTIFS(%s,%s,$B%d)" % (DATED, Dcol, r), "#,##0", "section.%s.posts" % name)
        rv.put(r, 4, "=SUMIFS(%s,%s,%s,$B%d)" % (G, DATED, Dcol, r), "#,##0", "section.%s.views" % name)
        rv.put(r, 5, '=IFERROR(AVERAGEIFS(%s,%s,%s,$B%d),"")' % (G, DATED, Dcol, r), "#,##0",
               "section.%s.avg" % name)
        rv.put(r, 6, '=IF(E{r}="","",RANK(E{r},$E${a}:$E${b},0))'.format(r=r, a=sec_first, b=sec_last), "0",
               "section.%s.rank" % name)
        r += 1
    bar_chart(ws, "bar", "Average views by song section", 5, 2, head, sec_first, sec_last, "I%d" % top)
    r = max(r + 1, top + CHART_ROWS + 1)

    # the three "best" rows, now that their tables exist
    def best(label_col, value_col, first, last):
        labels = "$%s$%d:$%s$%d" % (label_col, first, label_col, last)
        values = "$%s$%d:$%s$%d" % (value_col, first, value_col, last)
        return '=IFERROR(INDEX(%s,MATCH(MAX(%s),%s,0)),"")' % (labels, values, values)

    ws.cell(kpi_rows["best_platform"], 3).value = best("B", "E", plat_first, plat_last)
    ws.cell(kpi_rows["best_length"], 3).value = best("B", "D", len_first, len_last)
    ws.cell(kpi_rows["best_section"], 3).value = best("B", "E", sec_first, sec_last)
    return rv.cells


# -------------------------------------------------------------- README ---
README = [
    ("title", "Marketing review"),
    ("lead", "Your post log and review dashboard. Log each post on Posts; the Review sheet works everything out. "
             "Works in Excel (Mac or Windows), Apple Numbers and Google Sheets. A quick preview shows the gray "
             "columns, the Review sheet and its charts empty: open the file in the app (see Opening it, below)."),
    ("gap", None),
    ("head", "Start here"),
    ("text", "The rows on Posts and Calendar are made-up examples, so you can see how it works. When you start "
             "your own log, select Posts A2:M25 and press Delete, then Calendar A2:G7. Leave the gray columns alone."),
    ("gap", None),
    ("head", "Log a post"),
    ("step", ("1", "On Posts, go to the first empty row.")),
    ("step", ("2", "Type the date it went live, pick the platform and the song section, and paste the clip name "
                   "and the post's link.")),
    ("step", ("3", "Length (s): how long the clip is, in seconds.")),
    ("step", ("4", "About a week later, fill in Views, Likes, Comments, Shares and Saves from the app's insights. "
                   "Leave a box empty when the app doesn't show that number (YouTube shows no saves).")),
    ("step", ("5", "Trial reel: Yes for an Instagram trial reel, otherwise No.")),
    ("text", "A row counts as a post once it has a date: a row without one stays out of the Review sheet."),
    ("gap", None),
    ("head", "Clip names"),
    ("text", "Use the file name Take Cutter gave the clip, like chorus-2_01m22s-01m41s. The times say where the "
             "take sits in the long rehearsal recording, so you can always find it again."),
    ("gap", None),
    ("head", "Columns you type"),
    ("col", ("A  Date", "The day the post went live.")),
    ("col", ("B  Platform", "Instagram, TikTok or YouTube (pick from the list).")),
    ("col", ("C  Clip", "The Take Cutter file name.")),
    ("col", ("D  Section", "Which part of the song (pick from the list).")),
    ("col", ("E  Link", "The post's web address.")),
    ("col", ("F  Length (s)", "Clip length in seconds.")),
    ("col", ("G-K  Views to Saves", "From the app's insights. Plain numbers: 1200, not 1.2k.")),
    ("col", ("L  Trial reel", "Yes or No.")),
    ("col", ("M  Notes", "Anything: the time you posted, the title, what you tried.")),
    ("gap", None),
    ("head", "Formula columns (gray: don't type in them)"),
    ("col", ("N  Week", "Week number of the year; weeks start on Monday, and week 1 holds 1 January.")),
    ("col", ("O  Length bucket", "0–15 s, 15–25 s, 25–40 s or 40+ s.")),
    ("col", ("P  Engagement", "Likes + comments + shares + saves.")),
    ("col", ("Q  Engagement rate", "Engagement divided by views.")),
    ("col", ("R  Rank", "1 = most views, among the rows with a date. A tie goes to the post logged first.")),
    ("text", "They are filled down to row 1000, so a new row just works. If you type over one by accident, copy "
             "the gray cell above it and paste it back."),
    ("gap", None),
    ("head", "Length buckets"),
    ("col", ("0–15 s", "Under 15 seconds.")),
    ("col", ("15–25 s", "15 up to 25 seconds. A clip of exactly 15 s lands here.")),
    ("col", ("25–40 s", "25 up to 40 seconds.")),
    ("col", ("40+ s", "40 seconds and longer.")),
    ("text", "Change the cut-offs on the Lists sheet."),
    ("gap", None),
    ("head", "Titles"),
    ("text", "Shorts titles stop at 100 characters."),
    ("gap", None),
    ("head", "The other sheets"),
    ("col", ("Review", "Views by platform, average views by length, your top 10 clips with links, the last 12 "
                       "weeks, trial reels against main posts, and your best song sections. Nothing to type; "
                       "it updates as you log.")),
    ("col", ("Calendar", "Plan posts before they go out: date, song, section, the Take Cutter file, where it "
                         "goes, and a status (Idea, Cut, Edited, Scheduled, Posted). Once it is live, paste the "
                         "link and log it on Posts.")),
    ("col", ("Lists", "The dropdown choices and the length cut-offs. To change a choice, rename it there.")),
    ("gap", None),
    ("head", "Opening it"),
    ("col", ("Excel or Numbers", "Double-click the file. If Excel opens it in Protected View (a file from the "
                                 "internet), click Enable Editing.")),
    ("col", ("Google Sheets", "Upload the file to Google Drive, then Open with > Google Sheets.")),
    ("col", ("Quick previews", "The space bar in Finder, a phone's Files app and Excel's Protected View don't "
                               "calculate, so the gray columns, the Review sheet and its charts look empty there. "
                               "Open the file in the app and they fill in.")),
]


def build_readme(ws):
    set_widths(ws, {"A": 24, "B": 96})
    ws.sheet_view.showGridLines = False
    r = 1
    for kind, value in README:
        if kind == "title":
            ws.cell(r, 1, value).font = font(size=18, bold=True)
            ws.row_dimensions[r].height = 28
        elif kind == "lead":
            cell = ws.cell(r, 1, value)
            cell.font = font(size=11, color=INK2)
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
            cell.alignment = WRAP_TOP
            ws.row_dimensions[r].height = 15 * max(2, math.ceil(len(value) / 110))
        elif kind == "head":
            ws.cell(r, 1, value).font = font(size=12, bold=True)
        elif kind == "text":
            cell = ws.cell(r, 2, value)
            cell.font = font()
            cell.alignment = WRAP_TOP
            ws.row_dimensions[r].height = 14 * max(1, math.ceil(len(value) / 105))
        elif kind in ("step", "col"):
            label, text = value
            left = ws.cell(r, 1, label)
            left.font = font(bold=(kind == "col"), color=INK2)
            left.alignment = Alignment(vertical="top", horizontal="right" if kind == "step" else "left")
            cell = ws.cell(r, 2, text)
            cell.font = font()
            cell.alignment = WRAP_TOP
            ws.row_dimensions[r].height = 14 * max(1, math.ceil(len(text) / 105))
        r += 1


# ------------------------------------------------------------ workbook ---
def build_workbook(posts=SAMPLE_POSTS, calendar=SAMPLE_CALENDAR):
    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    posts_ws = wb.create_sheet("Posts")
    review = wb.create_sheet("Review")
    calendar_ws = wb.create_sheet("Calendar")
    lists = wb.create_sheet("Lists")
    build_readme(readme)
    build_lists(lists)
    build_posts(posts_ws, posts)
    cells = build_review(review)
    build_calendar(calendar_ws, calendar, posts)
    props = wb.properties
    props.creator = "Rainier"
    props.lastModifiedBy = "Rainier"
    props.title = "Marketing review"
    props.subject = "Post log and review dashboard (sample rows)"
    props.created = FIXED_TIME
    props.modified = FIXED_TIME
    return wb, cells


def workbook_bytes(wb):
    """Serialise without openpyxl's save() (which stamps 'modified' with the clock),
    then rewrite the zip with fixed entry times so the bytes are reproducible."""
    raw = io.BytesIO()
    ExcelWriter(wb, zipfile.ZipFile(raw, "w", zipfile.ZIP_DEFLATED, allowZip64=True)).save()
    src = zipfile.ZipFile(io.BytesIO(raw.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            entry = zipfile.ZipInfo(info.filename, date_time=ZIP_TIME)
            entry.compress_type = zipfile.ZIP_DEFLATED
            entry.create_system = 0
            entry.external_attr = 0
            dst.writestr(entry, src.read(info.filename), compresslevel=6)
    return out.getvalue()


# -------------------------------------------------------------- verify ---
def excel_serial(d):
    return float((d - datetime.date(1899, 12, 30)).days)


def weeknum_monday(d):
    """Excel WEEKNUM(d, 2): week 1 holds 1 January, weeks start on Monday."""
    jan1 = datetime.date(d.year, 1, 1)
    return (d.toordinal() - jan1.toordinal() + jan1.weekday()) // 7 + 1


def bucket_label(length):
    label = None
    for start, name in BUCKETS:
        if length >= start:
            label = name
    return label


def mean(xs):
    return sum(xs) / len(xs) if xs else ""


def ratio(a, b):
    return a / b if b else ""


# Made-up rows that exercise what the sample rows never do. --verify builds a workbook from
# each (in a temp folder, never the repo) and checks it like the committed one.
EDGE_POSTS = [
    # a tie (the row logged first ranks higher), bucket edges, a year boundary
    (D(2026, 12, 28), "TikTok", "tie-first", "chorus", 0, 5000, 400, 20, 30, 50, "No", ""),
    (D(2026, 12, 29), "Instagram", "tie-second", "verse 1", 14.99, 5000, 300, 10, 20, 40, "Yes", ""),
    # views typed before the date: more views than any other row, but not a post yet
    (None, "TikTok", "no-date-yet", "chorus", 15, 99000, 9000, 900, 900, 900, "No", "Date still to fill in"),
    (D(2026, 12, 30), "Instagram", "c04", "bridge", 15, 4200, 250, 9, 12, 30, None, "Trial reel left blank"),
    # a dated post with no platform, and one on a platform the list does not have
    (D(2026, 12, 31), None, "no-platform", None, 24.99, 700, 50, 2, 3, 4, None, ""),
    (D(2027, 1, 1), "Elsewhere", "unlisted-platform", "bridge", 25, 650, 40, 1, 2, 3, "No", ""),
    (D(2027, 1, 2), "YouTube", "zero-views", "intro", 39.99, 0, 0, 0, 0, None, "No", ""),
    (D(2027, 1, 3), "YouTube", "not-in-yet", "outro", 40, None, None, None, None, None, "No", ""),
    (D(2027, 1, 4), "TikTok", "c09", "chorus", 120, 3100, 200, 8, 10, 20, "No", ""),
    (D(2027, 1, 5), "TikTok", "c10", "pre-chorus", 22, 2900, 190, 7, 9, 18, "No", ""),
    (D(2027, 1, 5), "Instagram", "c11", "pre-chorus", 22, 2500, 170, 6, 8, 15, "Yes", ""),
    (D(2027, 1, 6), "YouTube", "c12", "verse 2", 30, 2300, 90, 3, 4, None, "No", ""),
    (D(2027, 1, 6), "TikTok", "no-length", "verse 2", None, 1900, 120, 5, 6, 10, "No", ""),
    # the week before the 12-week trend starts
    (D(2026, 10, 12), "TikTok", "before-the-trend", "intro", 18, 1500, 100, 4, 5, 8, "No", ""),
]

SCENARIOS = [("edge rows", EDGE_POSTS), ("no rows", [])]


def expected_values(posts):
    """The key numbers, computed straight from the rows without any spreadsheet logic.
    A row counts as a post once it has a date."""
    rows = []
    for i, p in enumerate(posts):
        date, platform, clip, section, length, views, likes, comments, shares, saves, trial, _notes = p
        eng = sum(v for v in (likes, comments, shares, saves) if v is not None)
        rows.append(dict(row=2 + i, date=date, platform=platform, clip=clip, section=section, length=length,
                         views=views, eng=eng, trial=trial, link=post_link(platform, i + 1)))
    dated_rows = [p for p in rows if p["date"] is not None]
    ranked = sorted((p for p in dated_rows if p["views"] is not None), key=lambda p: (-p["views"], p["row"]))
    for n, p in enumerate(ranked):
        p["rank"] = n + 1

    exp = {}
    # Posts: every row, plus blank rows further down
    for p in rows:
        r = p["row"]
        has_date = p["date"] is not None
        exp["Posts!N%d" % r] = weeknum_monday(p["date"]) if has_date else ""
        exp["Posts!O%d" % r] = bucket_label(p["length"]) if has_date and p["length"] is not None else ""
        exp["Posts!P%d" % r] = p["eng"] if has_date else ""
        exp["Posts!Q%d" % r] = ratio(p["eng"], p["views"] or 0) if has_date else ""
        exp["Posts!R%d" % r] = p.get("rank", "")
    for r in (2 + len(rows), 500, LAST):
        for col in "NOPQR":
            exp["Posts!%s%d" % (col, r)] = ""

    def agg(sel):
        sub = [p for p in dated_rows if sel(p)]
        vs = [p["views"] for p in sub if p["views"] is not None]
        eng = sum(p["eng"] for p in sub if (p["views"] or 0) > 0)
        return dict(posts=len(sub), views=sum(vs), avg=mean(vs), rate=ratio(eng, sum(vs)))

    def blank_if_zero(n):
        return n if n else ""

    total = agg(lambda p: True)
    r = {}
    r["kpi.posts"] = total["posts"]
    r["kpi.views"] = total["views"]
    r["kpi.avg"] = total["avg"]
    r["kpi.rate"] = total["rate"]
    plat = {name: agg(lambda p, n=name: p["platform"] == n) for name in PLATFORMS}
    for name, a in plat.items():
        for k in ("posts", "views", "avg", "rate"):
            r["platform.%s.%s" % (name, k)] = a[k]
        r["platform.%s.share" % name] = ratio(a["views"], total["views"])
    other = agg(lambda p: p["platform"] not in PLATFORMS)
    has_other = other["posts"] > 0
    r["platform.other.label"] = OTHER_PLATFORM if has_other else ""
    r["platform.other.posts"] = blank_if_zero(other["posts"])
    r["platform.other.views"] = other["views"] if has_other else ""
    r["platform.other.share"] = ratio(other["views"], total["views"]) if has_other else ""
    r["platform.all.posts"] = total["posts"]
    r["platform.all.views"] = total["views"]
    r["platform.all.avg"] = total["avg"]
    r["platform.all.rate"] = total["rate"]
    r["platform.all.share"] = ratio(total["views"], total["views"])
    lens = {}
    for i, (_start, label) in enumerate(BUCKETS):
        a = agg(lambda p, lb=label: p["length"] is not None and bucket_label(p["length"]) == lb)
        lens[label] = a
        r["length.%d.posts" % i] = a["posts"]
        r["length.%d.avg" % i] = a["avg"]
        r["length.%d.rate" % i] = a["rate"]
    for k in range(1, 11):
        p = ranked[k - 1] if k <= len(ranked) else None
        r["top.%d.clip" % k] = p["clip"] if p else ""
        r["top.%d.platform" % k] = (p["platform"] or "") if p else ""
        r["top.%d.section" % k] = (p["section"] or "") if p else ""
        r["top.%d.length" % k] = (p["length"] if p["length"] is not None else "") if p else ""
        r["top.%d.views" % k] = p["views"] if p else ""
        r["top.%d.link" % k] = p["link"] if p else ""
    if dated_rows:
        latest = max(p["date"] for p in dated_rows)
        first = latest - datetime.timedelta(days=latest.weekday() + 77)
    for k in range(12):
        if not dated_rows:
            for name in ("start", "posts", "views", "avg"):
                r["trend.%d.%s" % (k, name)] = ""
            continue
        start = first + datetime.timedelta(days=7 * k)
        end = start + datetime.timedelta(days=7)
        a = agg(lambda p, s=start, e=end: s <= p["date"] < e)
        r["trend.%d.start" % k] = excel_serial(start)
        r["trend.%d.posts" % k] = a["posts"]
        r["trend.%d.views" % k] = a["views"]
        r["trend.%d.avg" % k] = a["avg"]
    for key, sel in (("main", lambda p: p["trial"] != "Yes"),
                     ("ig_main", lambda p: p["platform"] == "Instagram" and p["trial"] != "Yes"),
                     ("trial", lambda p: p["trial"] == "Yes")):
        a = agg(sel)
        for k in ("posts", "views", "avg", "rate"):
            r["trial.%s.%s" % (key, k)] = a[k]
    secs = {name: agg(lambda p, n=name: p["section"] == n) for name in SECTIONS}
    avgs = sorted((a["avg"] for a in secs.values() if a["avg"] != ""), reverse=True)
    for name, a in secs.items():
        r["section.%s.posts" % name] = a["posts"]
        r["section.%s.views" % name] = a["views"]
        r["section.%s.avg" % name] = a["avg"]
        r["section.%s.rank" % name] = (avgs.index(a["avg"]) + 1) if a["avg"] != "" else ""

    def best(groups):
        scored = [(a["avg"], i, name) for i, (name, a) in enumerate(groups) if a["avg"] != ""]
        return max(scored, key=lambda t: (t[0], -t[1]))[2] if scored else ""

    r["kpi.best_platform"] = best(list(plat.items()))
    r["kpi.best_length"] = best(list(lens.items()))
    r["kpi.best_section"] = best(list(secs.items()))
    return exp, r


def check_workbook(path, posts, label):
    """Evaluate every formula in the workbook at path; compare the key numbers with
    expected_values(posts). Returns the number of problems."""
    import formulas  # optional: only needed for --verify
    from formulas.tokens.operand import XlError

    _wb, cells = build_workbook(posts, [])
    exp_posts, exp_review = expected_values(posts)
    expected = dict(exp_posts)
    for name, value in exp_review.items():
        expected["Review!" + cells[name]] = value

    book = load_workbook(path)
    formula_cells = []
    for ws in book.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_cells.append("%s!%s" % (ws.title, cell.coordinate))

    model = formulas.ExcelModel().loads(path).finish()
    solution = model.calculate()
    values = {}
    for key, ranges in solution.items():
        text = str(key)  # formulas 1.3.x keys look like '[file]SHEET'!A1
        if not text.startswith("'[") or ":" in text.split("!")[-1]:
            continue
        book_sheet, coord = text.rsplit("!", 1)
        sheet = book_sheet.split("]", 1)[1].rstrip("'")
        values[(sheet.upper(), coord.upper())] = ranges.value[0][0] if hasattr(ranges, "value") else ranges

    def value_at(ref):
        sheet, coord = ref.split("!")
        return values.get((sheet.upper(), coord.upper()), "<missing>")

    errors, missing = [], []
    for ref in formula_cells:
        v = value_at(ref)
        if isinstance(v, str) and v == "<missing>":
            missing.append(ref)
        elif isinstance(v, XlError):
            errors.append("%s = %s" % (ref, v))

    mismatches = []
    for ref, want in sorted(expected.items()):
        got = value_at(ref)
        if isinstance(got, XlError):
            mismatches.append("%s: error %s, want %r" % (ref, got, want))
            continue
        if want == "" or isinstance(want, str):
            ok = (got == want) or (want == "" and got is None)
        else:
            try:
                ok = math.isclose(float(got), float(want), rel_tol=1e-9, abs_tol=1e-9)
            except (TypeError, ValueError):
                ok = False
        if not ok:
            mismatches.append("%s: got %r, want %r" % (ref, got, want))

    if not formula_cells:
        missing.append("(no formula cells found)")
    print("verify [%s]: %d formula cells evaluated, %d error values, %d not evaluated" % (
        label, len(formula_cells), len(errors), len(missing)))
    print("verify [%s]: %d key values compared with an independent computation, %d mismatches" % (
        label, len(expected), len(mismatches)))
    for line in (errors + missing + mismatches)[:40]:
        print("  " + line)
    for name in ("kpi.posts", "kpi.views", "kpi.avg", "kpi.rate", "kpi.best_platform", "kpi.best_length",
                 "kpi.best_section", "top.1.clip", "top.1.views"):
        print("  %-18s %s" % (name, value_at("Review!" + cells[name])))
    return len(errors) + len(missing) + len(mismatches)


def verify(path):
    """Evaluate every formula with the 'formulas' package, in the workbook at path (the
    sample rows) and in a scratch workbook per scenario; fail on any error value or on
    any key number that differs from expected_values()."""
    try:
        import formulas  # optional: only needed for --verify
    except ImportError:
        print("verify: the 'formulas' package is not installed (pip install formulas==%s)" % VERIFIED_WITH,
              file=sys.stderr)
        return 2
    if formulas.__version__ != VERIFIED_WITH:
        print("note: formulas %s (--verify was written against %s and reads its internal cell keys)" % (
            formulas.__version__, VERIFIED_WITH), file=sys.stderr)

    def hyperlink(link, name=None):  # formulas has no HYPERLINK; Excel shows the name, else the link
        return link if name is None else name

    formulas.get_functions()["HYPERLINK"] = hyperlink

    problems = check_workbook(path, SAMPLE_POSTS, "sample rows: %s" % os.path.basename(path))
    with tempfile.TemporaryDirectory(prefix="marketing-review-verify-") as tmp:
        for label, posts in SCENARIOS:
            scratch = os.path.join(tmp, "scenario.xlsx")
            wb, _cells = build_workbook(posts, [])
            with open(scratch, "wb") as fh:
                fh.write(workbook_bytes(wb))
            problems += check_workbook(scratch, posts, label)
    print("verify: %s" % ("PASS" if problems == 0 else "FAIL, %d problems" % problems))
    return 1 if problems else 0


def main(argv):
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--out", default=DEFAULT_OUT, help="where to write the workbook (default: kits/marketing-review.xlsx)")
    ap.add_argument("--verify", nargs="?", const=DEFAULT_OUT, metavar="XLSX",
                    help="evaluate every formula in XLSX (default: the committed workbook) instead of building")
    args = ap.parse_args(argv)
    if args.verify:
        return verify(args.verify)
    if openpyxl.__version__ != BUILT_WITH:
        print("note: openpyxl %s (the committed file was built with %s); the bytes may differ" % (
            openpyxl.__version__, BUILT_WITH), file=sys.stderr)
    wb, _cells = build_workbook()
    data = workbook_bytes(wb)
    with open(args.out, "wb") as fh:
        fh.write(data)
    print("wrote %s (%d bytes)" % (args.out, len(data)))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
